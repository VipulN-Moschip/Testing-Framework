
import pytest
import openpyxl
from datetime import datetime
from configuration import TESTRAIL_URL, USERNAME, PASSWORD, TEST_IDS, TEST_CASE_IDS
from testrail_api import TestRailAPI

# Initialize TestRail API
test_rail_api = TestRailAPI(TESTRAIL_URL, USERNAME, PASSWORD)

#  List to store test results
test_results = []

def update_testrail(test_id, status, comment=""):
    """ Updates a test result in TestRail using the test ID. """
    try:
        response = test_rail_api.results.add_result(test_id, status_id=status, comment=comment)
        print(f" TestRail Updated: Test {test_id} -> Status {status}")
        return response
    except Exception as e:
        print(f" Failed to update TestRail for test {test_id}: {str(e)}")
        return None

@pytest.hookimpl(tryfirst=True, hookwrapper=True)
def pytest_runtest_makereport(item, call):
    """ Pytest hook to collect test results for Excel reporting and update TestRail. """
    outcome = yield
    report = outcome.get_result()

    testrail_marker = item.get_closest_marker("testrail")
    if testrail_marker:
        test_run_id = testrail_marker.args[0]
        test_case_id = TEST_CASE_IDS.get(test_run_id, test_run_id)

        if report.when == "call":
            status = 1 if report.passed else 5  # 1 = Passed, 5 = Failed
            comment = "Test passed successfully" if report.passed else f"Test failed: {report.longrepr}"

            #  Fetch test details from TestRail
            test_details = fetch_testrail_case(test_case_id)

            #  Store test results for Excel
            test_results.append({
                "Test Case ID": test_case_id,
                "Test Case Name": test_details["Test Case Name"],
                "Steps": test_details["Steps"],
                "Expected Result": test_details["Expected Result"],
                "Actual Result": "Passed" if report.passed else "Failed",
                "Status": "✔" if report.passed else "❌",
                "Execution Duration": f"{report.duration:.3f} sec",  #  Use Pytest duration
                "Timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            })

            #  Update TestRail
            update_testrail(test_run_id, status, comment)

def fetch_testrail_case(case_id):
    """ Fetches test case details from TestRail using the correct case ID. """
    try:
        response = test_rail_api.cases.get_case(case_id)
        return {
            "Test Case ID": case_id,
            "Test Case Name": response.get("title", f"Unknown (ID {case_id})"),
            "Steps": response.get("custom_steps", "No steps provided"),
            "Expected Result": response.get("custom_expected", "No expected result provided"),
        }
    except Exception as e:
        print(f" Failed to fetch TestRail case {case_id}: {str(e)}")
        return {
            "Test Case ID": case_id,
            "Test Case Name": f"Unknown (ID {case_id})",
            "Steps": "Failed to fetch steps",
            "Expected Result": "Failed to fetch expected result",
        }

@pytest.fixture(scope="session", autouse=True)
def generate_excel_report():
    """ Generates an Excel report after all tests finish. """
    yield  #  Wait for test execution to complete

    #  Create Excel workbook
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Test Results"

    #  Define Headers
    headers = [
        "Test Case ID", "Test Case Name", "Steps", "Expected Result",
        "Actual Result", "Status", "Execution Duration", "Timestamp"
    ]
    sheet.append(headers)

    #  Apply Formatting: Wrap Text and Adjust Column Width
    for col_num, header in enumerate(headers, 1):
        col_letter = sheet.cell(row=1, column=col_num).column_letter
        sheet.column_dimensions[col_letter].width = 20  #  Set column width
        sheet.cell(row=1, column=col_num).alignment = openpyxl.styles.Alignment(wrap_text=True)

    #  Write test results with formatting
    for row_idx, result in enumerate(test_results, start=2):
        row_data = list(result.values())
        sheet.append(row_data)

        #  Apply Wrap Text for all cells
        for col_idx in range(1, len(headers) + 1):
            sheet.cell(row=row_idx, column=col_idx).alignment = openpyxl.styles.Alignment(wrap_text=True)

        # Apply coloring for Pass/Fail
        status_cell = sheet.cell(row=row_idx, column=6)
        if result["Actual Result"] == "Passed":
            status_cell.fill = openpyxl.styles.PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Green
        else:
            status_cell.fill = openpyxl.styles.PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Red

    #  Save Excel Report
    excel_path = "reports/test_results.xlsx"
    workbook.save(excel_path)
    print(f"Excel report saved at: {excel_path}")
